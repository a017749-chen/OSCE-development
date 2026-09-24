"""The builder produces a station the validator passes, and the validator catches
the problems it claims to catch."""
import copy

import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "scripts"))


import build_station  # noqa: E402
import osce_rules  # noqa: E402
import validate_station  # noqa: E402

RULES = osce_rules.load()
EXAMPLE = build_station.load_station(REPO / "examples" / "example_station.yaml")


class StationToolTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def build(self, station, name="station.docx"):
        path = Path(self.tmp.name) / name
        build_station.build(station, RULES, REPO / "templates").save(str(path))
        return path

    def results(self, path):
        report = validate_station.validate(path, RULES)
        return {check: (status, detail) for status, check, detail in report.results}, report

    def test_the_example_builds_and_passes_every_check(self):
        results, report = self.results(self.build(EXAMPLE))
        failed = {k: v for k, v in results.items() if v[0] in ("FAIL", "WARN")}
        self.assertEqual(failed, {})
        self.assertEqual(report.station_type, "explanation")

    def test_score_total_follows_the_item_count(self):
        station = copy.deepcopy(EXAMPLE)
        station["checklist"] = station["checklist"][:10]
        results, _ = self.results(self.build(station))
        self.assertEqual(results["滿分＝N×2"], ("PASS", "20 分"))

    def test_the_builder_refuses_content_that_breaks_a_hard_rule(self):
        cases = {
            "sign": lambda s: s.update(sign="一" * 31),
            "tasks": lambda s: s["candidate"].update(tasks=["甲", "乙", "丙", "丁"]),
            "questions": lambda s: [row.update(question=True) for row in s["sp"]["dialogue"]],
            "stars": lambda s: [item.pop("star", None) for item in s["checklist"]],
            "differentials": lambda s: s["examiner"].update(differentials=["一", "二"]),
            "teach-back": lambda s: s["sp"].update(dialogue=s["sp"]["dialogue"][:-1]),
        }
        for name, spoil in cases.items():
            station = copy.deepcopy(EXAMPLE)
            spoil(station)
            with self.subTest(name), self.assertRaises(build_station.StationError):
                build_station.build(station, RULES, REPO / "templates")

    def test_the_builder_never_overwrites(self):
        target = Path(self.tmp.name) / "existing.docx"
        target.write_bytes(b"hand-edited")
        code = build_station.main([str(REPO / "examples" / "example_station.yaml"), "-o", str(target)])
        self.assertEqual(code, 2)
        self.assertEqual(target.read_bytes(), b"hand-edited")

    def test_the_validator_catches_what_it_claims_to(self):
        import docx
        from docx.oxml.ns import qn
        from docx.shared import Pt

        path = self.build(EXAMPLE)
        document = docx.Document(str(path))
        # Too many SP questions, wrong total, shrunken sign, borders stripped, PS 5.1 residue.
        dialogue = document.tables[-2]
        for row in dialogue.rows[1:]:
            row.cells[2].text = "那怎麼辦？"
        for p in document.paragraphs:
            if p.text.startswith("滿分"):
                p.runs[0].text = "滿分：30 分"
        document.tables[0].rows[0].cells[0].paragraphs[0].runs[0].font.size = Pt(36)
        border = document.tables[1]._tbl.tblPr.find(qn("w:tblBorders"))
        document.tables[1]._tbl.tblPr.remove(border)
        document.add_paragraph("血壓 120/80u3000mmHg")
        broken = Path(self.tmp.name) / "broken.docx"
        document.save(str(broken))

        results, report = self.results(broken)
        self.assertTrue(report.failed)
        self.assertEqual(results["對白表提問列 ≤2"][0], "FAIL")
        self.assertEqual(results["滿分＝N×2"][0], "FAIL")
        self.assertEqual(results["告示牌站次號 48pt"][0], "FAIL")
        self.assertEqual(results["表格有框線"][0], "FAIL")
        self.assertEqual(results["無殘留 Unicode 逸出字面"][0], "FAIL")

    def test_the_dialogue_header_follows_the_official_template_per_station_type(self):
        station = copy.deepcopy(EXAMPLE)
        station["station_type"] = "history"
        document = build_station.build(station, RULES, REPO / "templates")
        header = document.tables[-2].rows[0].cells[2].text
        self.assertEqual(header, RULES["station_types"]["history"]["dialogue_third_column"])

    def test_narrative_rubric_builds(self):
        station = copy.deepcopy(EXAMPLE)
        station["rubric"] = "narrative"
        station["narrative"] = {d["name"]: {"4": "優", "3": "熟", "2": "展", "1": "未"}
                                for d in RULES["narrative"]["scored_dimensions"]}
        results, _ = self.results(self.build(station))
        self.assertEqual(results["條列式評分表"][0], "SKIP")
        self.assertEqual(results["表格有框線"][0], "PASS")


if __name__ == "__main__":
    unittest.main()
