"""Check an OSCE station .docx against rules.yaml.

    python scripts/validate_station.py stations/某站.docx [更多檔案...]
    python scripts/validate_station.py --report stations/STATUS.md stations/*.docx

Checks what can be checked mechanically - page, fonts, section order, sign and
background length, red bullets, score total against item count, star items, SP
question rows, the official dialogue header, teach-back, table widths, and escape
sequences that PowerShell 5.1 leaves behind. It does not judge clinical content,
anchors or realism; a station that passes still needs a human review and a page-by-
page look at the rendered document.

Works on files built by build_station.py and on older hand-made or Word-COM files:
font sizes are resolved the way Word resolves them (run, then paragraph style, then
style inheritance, then document default).
"""
from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import docx
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
import osce_rules  # noqa: E402

PASS, WARN, FAIL, SKIP = "PASS", "WARN", "FAIL", "SKIP"
SECTIONS = ["告示牌", "考生指引", "評分表", "考官指引", "SP指引"]
# PowerShell 5.1 prints `u3000 as the literal "u3000", usually right after a number or a
# CJK character ("120/80u3000mmHg"), so only a preceding Latin letter rules a match out.
ESCAPES = re.compile(r"(?<![A-Za-z])u(3000|2610|25A0|25A1|2103|03BC|2019|2026)(?![0-9A-Fa-f])")
QUESTION = re.compile(r"[？?]\s*[」』\"]?\s*$")


@dataclass
class Report:
    path: Path
    station_type: str | None = None
    results: list = field(default_factory=list)

    def add(self, status, check, detail=""):
        self.results.append((status, check, detail))

    @property
    def failed(self):
        return any(s == FAIL for s, _, _ in self.results)

    def counts(self):
        return {s: sum(1 for r in self.results if r[0] == s) for s in (PASS, WARN, FAIL, SKIP)}


def plain(text: str) -> str:
    """Compare headings without numbering, spaces or full-width variants."""
    text = unicodedata.normalize("NFKC", text)
    text = re.sub(r"^[一二三四五六七八九十]+[、.．]\s*", "", text.strip())
    return re.sub(r"[\s　()（）劇本]", "", text)


def chars(text: str) -> int:
    return len(re.sub(r"[\s　]", "", text))


def default_size(document) -> float | None:
    styles = document.styles.element
    sz = styles.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}/{qn('w:sz')}")
    return int(sz.get(qn("w:val"))) / 2 if sz is not None else None


def style_size(style) -> float | None:
    while style is not None:
        if style.font is not None and style.font.size is not None:
            return style.font.size.pt
        style = style.base_style
    return None


def size_of(run, paragraph, fallback) -> float | None:
    if run.font.size is not None:
        return run.font.size.pt
    return style_size(paragraph.style) or fallback


def paragraph_sizes(paragraph, fallback) -> set:
    return {size_of(r, paragraph, fallback) for r in paragraph.runs if r.text.strip()}


def is_red(run) -> bool:
    color = run.font.color
    return color is not None and color.type is not None and color.rgb is not None and str(color.rgb) in {"FF0000", "C00000"}


def table_width_cm(table) -> float | None:
    cells = table.rows[0]._tr.findall(qn("w:tc"))
    total = 0
    for tc in cells:
        pr = tc.tcPr
        w = pr.find(qn("w:tcW")) if pr is not None else None
        if w is None or w.get(qn("w:type")) != "dxa":
            return None
        total += int(w.get(qn("w:w")))
    return total / 567


def detect_type(document, path: Path, rules) -> str | None:
    """Station type from a checked box, then from the filename."""
    names = {key: v["name"] for key, v in rules["station_types"].items()}
    for p in document.paragraphs[:120]:
        for key, name in names.items():
            if re.search(rf"[■☑✓]\s*{name}", p.text):
                return key
    for key, name in names.items():
        short = name[:4]
        if short in path.stem or name in path.stem:
            return key
    return None


def validate(path: Path, rules: dict) -> Report:
    report = Report(path)
    document = docx.Document(str(path))
    fallback = default_size(document)
    fonts = rules["fonts_pt"]
    limits = rules["limits"]
    paragraphs = document.paragraphs

    # Page
    section = document.sections[0]
    a4 = abs(section.page_width.cm - 21.0) < 0.2 and abs(section.page_height.cm - 29.7) < 0.2
    margins = [section.left_margin, section.right_margin, section.top_margin, section.bottom_margin]
    ok_margins = all(m is not None and abs(m.cm - rules["page"]["margin_cm"]) < 0.15 for m in margins)
    report.add(PASS if a4 else FAIL, "A4 紙張", f"{section.page_width.cm:.1f}×{section.page_height.cm:.1f} cm")
    report.add(PASS if ok_margins else WARN, f"邊界 {rules['page']['margin_cm']} cm",
               "、".join(f"{m.cm:.2f}" for m in margins if m is not None))

    # Section order
    found = []
    for i, p in enumerate(paragraphs):
        key = plain(p.text)
        for name in SECTIONS:
            if key == plain(name) and name not in [n for n, _ in found]:
                found.append((name, i))
    names_found = [n for n, _ in found]
    missing = [n for n in SECTIONS if n not in names_found]
    in_order = names_found == [n for n in SECTIONS if n in names_found]
    report.add(FAIL if missing or not in_order else PASS, "五大部分齊全且依序",
               ("缺：" + "、".join(missing)) if missing else ("順序不對" if not in_order else ""))
    at = dict(found)
    heading_sizes = {n: paragraph_sizes(paragraphs[i], fallback) for n, i in found}
    wrong = [f"{n} {sorted(s)}" for n, s in heading_sizes.items() if s and s != {fonts['section_title']}]
    report.add(FAIL if wrong else PASS, f"部分標題 {fonts['section_title']}pt", "；".join(wrong))

    report.station_type = detect_type(document, path, rules)
    stype = rules["station_types"].get(report.station_type) if report.station_type else None
    report.add(PASS if stype else WARN, "站型可辨識", stype["name"] if stype else "未勾選站型，檔名也無法判斷")

    tables = document.tables

    # Sign: the first table, two rows, one column (Word-COM and builder layout)
    sign = next((t for t in tables if len(t.columns) == 1 and len(t.rows) >= 2), None)
    if sign is None:
        report.add(SKIP, "告示牌字級與字數", "告示牌不是表格（可能在文字方塊內）")
    else:
        number, info = sign.rows[0].cells[0], sign.rows[1].cells[0]
        n_sizes = {s for p in number.paragraphs for s in paragraph_sizes(p, fallback)}
        i_sizes = {s for p in info.paragraphs for s in paragraph_sizes(p, fallback)}
        report.add(PASS if n_sizes == {fonts["sign_station_number"]} else FAIL,
                   f"告示牌站次號 {fonts['sign_station_number']}pt", str(sorted(n_sizes)))
        report.add(PASS if i_sizes == {fonts["sign_patient_info"]} else FAIL,
                   f"告示牌病患資訊 {fonts['sign_patient_info']}pt", str(sorted(i_sizes)))
        length = chars(info.text)
        report.add(PASS if 0 < length <= limits["sign_patient_info_chars"] else FAIL,
                   f"告示牌病患資訊 ≤{limits['sign_patient_info_chars']} 字",
                   f"{length} 字" if length else "空白")

    # Candidate guide
    if "考生指引" in at and "評分表" in at:
        block = paragraphs[at["考生指引"] + 1:at["評分表"]]
        body = [p for p in block if re.match(r"^[■●]?\s*(背景資料|測驗主題|測驗時間|相關檢查報告)", p.text.strip())]
        sizes = {s for p in body for s in paragraph_sizes(p, fallback)}
        report.add(PASS if sizes == {fonts["candidate_body"]} else FAIL,
                   f"考生指引本體 {fonts['candidate_body']}pt", str(sorted(sizes)) if sizes else "找不到欄位")
        bullets = [p for p in block if p.text.strip().startswith("●")]
        red = [p for p in bullets if any(is_red(r) for r in p.runs if r.text.strip())]
        report.add(PASS if len(bullets) <= limits["candidate_tasks_max"] else FAIL,
                   f"測驗主題 ≤{limits['candidate_tasks_max']} 個 ●", f"{len(bullets)} 個")
        report.add(PASS if bullets and len(red) == len(bullets) else WARN, "測驗主題為紅色",
                   f"{len(red)}/{len(bullets)} 個是紅色")
        background = next((p.text for p in block if "背景資料" in p.text), "")
        text = re.sub(r"^[■\s]*背景資料[:：]\s*", "", background.strip())
        length = chars(text)
        report.add(PASS if 0 < length <= limits["candidate_background_chars"] else FAIL,
                   f"背景資料 ≤{limits['candidate_background_chars']} 字", f"{length} 字")
    else:
        report.add(SKIP, "考生指引", "找不到考生指引或評分表標題")

    # Checklist
    checklist = next((t for t in tables if t.rows[0].cells[0].text.strip().startswith("評分項目")), None)
    if checklist is None:
        report.add(SKIP, "條列式評分表", "沒有「評分項目」表（可能是敘事醫學評分）")
    else:
        items = [r for r in checklist.rows[1:] if r.cells[0].text.strip()]
        n = len(items)
        stars = sum(1 for r in items if "★" in r.cells[0].text)
        c = rules["checklist"]
        report.add(PASS if c["items_min"] <= n <= c["items_max"] else WARN,
                   f"評分項目 {c['items_min']}–{c['items_max']} 項", f"{n} 項")
        stated = None
        for p in paragraphs:
            m = re.search(r"滿分[：:]\s*(\d+)\s*分", p.text)
            if m:
                stated = int(m.group(1))
                break
        report.add(PASS if stated == 2 * n else FAIL, "滿分＝N×2",
                   f"寫 {stated} 分，{n} 項應為 {2 * n} 分" if stated != 2 * n else f"{stated} 分")
        report.add(PASS if c["high_discrimination_min"] <= stars <= c["high_discrimination_max"] else WARN,
                   f"★ {c['high_discrimination_min']}–{c['high_discrimination_max']} 項",
                   f"{stars} 項" + ("（未標★）" if stars == 0 else ""))

    # Dialogue table
    # Found by its SP columns, not its first header: older stations renamed that column.
    dialogue = next((t for t in tables if len(t.columns) == 3
                     and "SP" in unicodedata.normalize("NFKC", t.rows[0].cells[1].text)
                     and "SP" in unicodedata.normalize("NFKC", t.rows[0].cells[2].text)), None)
    if dialogue is None:
        report.add(FAIL, "SP 三欄對白表", "找不到三欄對白表")
    else:
        first = dialogue.rows[0].cells[0].text.strip()
        official = rules["sp"]["dialogue_columns"][0]
        report.add(PASS if plain(first) == plain(official) else WARN, "對白表第一欄標題依官方範本",
                   "" if plain(first) == plain(official) else f"寫「{first}」，官方為「{official}」")
        header = plain(dialogue.rows[0].cells[2].text)
        if stype:
            expected = plain(stype["dialogue_third_column"])
            report.add(PASS if header == expected else WARN, "對白表第三欄標題依官方範本",
                       f"寫「{dialogue.rows[0].cells[2].text.strip()}」，官方為「{stype['dialogue_third_column']}」"
                       if header != expected else "")
        questions = sum(1 for r in dialogue.rows[1:] if QUESTION.search(r.cells[2].text.strip()))
        cap = rules["sp"]["dialogue_question_rows_max"]
        report.add(PASS if questions <= cap else FAIL, f"對白表提問列 ≤{cap}", f"{questions} 列以問號結尾")

    # Teach-back
    if stype and stype.get("teach_back_required"):
        text = "\n".join(p.text for p in paragraphs) + "\n".join(c.text for t in tables for r in t.rows for c in r.cells)
        has = re.search(r"teach[\s-]?back|回述|複述|請病人.*說一次", text, re.I)
        report.add(PASS if has else FAIL, "Teach-back", "" if has else "全文找不到 Teach-back／回述")

    # Table widths
    over = []
    for i, t in enumerate(tables):
        w = table_width_cm(t)
        if w is not None and w > rules["page"]["table_width_max_cm"] + 0.05:
            over.append(f"表 {i + 1}：{w:.1f} cm")
    report.add(FAIL if over else PASS, f"表格總寬 ≤{rules['page']['table_width_max_cm']} cm", "；".join(over))

    # Borders: a table with neither its own borders nor a bordered style prints as loose text.
    def bordered(t):
        own = t._tbl.tblPr.find(qn("w:tblBorders")) if t._tbl.tblPr is not None else None
        if own is not None and any(e.get(qn("w:val")) not in (None, "nil", "none") for e in own):
            return True
        style = t.style
        while style is not None:
            b = style.element.find(f"{qn('w:tblPr')}/{qn('w:tblBorders')}")
            if b is not None and any(e.get(qn("w:val")) not in (None, "nil", "none") for e in b):
                return True
            style = style.base_style
        cell_borders = t._tbl.findall(f".//{qn('w:tcBorders')}")
        return bool(cell_borders)
    bare = [str(i + 1) for i, t in enumerate(tables) if not bordered(t)]
    report.add(FAIL if bare else PASS, "表格有框線", ("無框線：表 " + "、".join(bare)) if bare else "")

    # PowerShell 5.1 escape residue
    full = "\n".join(p.text for p in paragraphs) + "\n".join(c.text for t in tables for r in t.rows for c in r.cells)
    residue = sorted(set(m.group(0) for m in ESCAPES.finditer(full)))
    report.add(FAIL if residue else PASS, "無殘留 Unicode 逸出字面", "、".join(residue))
    return report


def is_font(check: str) -> bool:
    return "pt" in check


def render_markdown(reports) -> str:
    lines = ["# 教案檢核報告", "",
             "由 `scripts/validate_station.py` 產生，只檢查機械可驗證的格式規則；**通過不代表臨床內容或版面已審閱**。",
             "教案檔案未被修改。重新產生：`python scripts/validate_station.py --report stations/STATUS.md stations/*.docx`", ""]
    font_debt = [r for r in reports if any(s == FAIL and is_font(n) for s, n, _ in r.results)]
    if font_debt:
        lines += [f"**字級：{len(font_debt)}/{len(reports)} 份仍是舊字級**（標題、告示牌、考生指引未達官方 "
                  "20／48／36／26 pt）。下表先列字級以外的問題。", ""]
    lines += ["| 教案 | 站型 | 失敗 | 警告 | 字級以外的問題 |", "|---|---|---|---|---|"]
    for r in reports:
        c = r.counts()
        problems = [f"{name}（{detail}）" if detail else name
                    for s, name, detail in r.results if s in (FAIL, WARN) and not is_font(name)]
        stype = r.station_type or "—"
        lines.append(f"| {r.path.name} | {stype} | {c[FAIL]} | {c[WARN]} | {'；'.join(problems) or '—'} |")
    for r in reports:
        lines += ["", f"## {r.path.name}", ""]
        lines += [f"- **{s}** {name}" + (f"：{detail}" if detail else "") for s, name, detail in r.results]
    return "\n".join(lines) + "\n"


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("files", nargs="+")
    parser.add_argument("--report", help="write a Markdown report here instead of printing")
    args = parser.parse_args(argv)
    rules = osce_rules.load()
    reports = [validate(Path(f), rules) for f in args.files if not Path(f).name.startswith("~$")]
    if args.report:
        Path(args.report).write_text(render_markdown(reports), encoding="utf-8", newline="\n")
        print(f"已寫入 {args.report}：{len(reports)} 份，{sum(r.failed for r in reports)} 份有失敗項。")
    else:
        for r in reports:
            print(f"\n{r.path.name}  [{r.station_type or '?'}]")
            for s, name, detail in r.results:
                print(f"  {s:4}  {name}" + (f"：{detail}" if detail else ""))
    return 1 if any(r.failed for r in reports) else 0


if __name__ == "__main__":
    sys.exit(main())
